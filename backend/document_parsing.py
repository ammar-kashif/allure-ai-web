"""Slide and document parsing for PDF and PPTX files.

Extracts structured content (title, bullets, notes) per page/slide.
"""

import logging
import os
from typing import Any

logger = logging.getLogger(__name__)


def _parse_pdf(file_path: str) -> list[dict[str, Any]]:
    """Extract text per page from a PDF using pdfplumber."""
    try:
        import pdfplumber  # type: ignore
    except ImportError:
        raise ImportError("pdfplumber is required for PDF parsing. Run: pip install pdfplumber")

    slides: list[dict[str, Any]] = []
    with pdfplumber.open(file_path) as pdf:
        for i, page in enumerate(pdf.pages):
            text = page.extract_text() or ""
            lines = [ln.strip() for ln in text.splitlines() if ln.strip()]
            title = lines[0] if lines else f"Page {i + 1}"
            content = "\n".join(lines[1:]) if len(lines) > 1 else ""
            slides.append(
                {
                    "index": i,
                    "title": title,
                    "content": content,
                    "notes": "",
                }
            )
    return slides


def _parse_pptx(file_path: str) -> list[dict[str, Any]]:
    """Extract title, bullets, and speaker notes per slide from a PPTX."""
    try:
        from pptx import Presentation  # type: ignore
        from pptx.util import Pt  # type: ignore  # noqa: F401
    except ImportError:
        raise ImportError("python-pptx is required for PPTX parsing. Run: pip install python-pptx")

    prs = Presentation(file_path)
    slides: list[dict[str, Any]] = []

    for i, slide in enumerate(prs.slides):
        title = ""
        bullets: list[str] = []
        notes = ""

        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            # Title placeholder
            if shape.shape_id == 1 or (hasattr(shape, "placeholder_format") and shape.placeholder_format and shape.placeholder_format.idx == 0):
                title = shape.text_frame.text.strip()
            else:
                for para in shape.text_frame.paragraphs:
                    text = para.text.strip()
                    if text:
                        bullets.append(text)

        # Speaker notes
        if slide.has_notes_slide:
            notes_frame = slide.notes_slide.notes_text_frame
            if notes_frame:
                notes = notes_frame.text.strip()

        slides.append(
            {
                "index": i,
                "title": title or f"Slide {i + 1}",
                "content": "\n".join(bullets),
                "notes": notes,
            }
        )
    return slides


def parse_document(file_path: str) -> dict[str, Any]:
    """Parse a PDF or PPTX file and return structured slide content.

    Returns:
        {
          "filename": str,
          "type": "pdf" | "pptx",
          "slides": [ { index, title, content, notes } ]
        }
    """
    ext = os.path.splitext(file_path)[1].lower()
    filename = os.path.basename(file_path)

    if ext == ".pdf":
        slides = _parse_pdf(file_path)
        doc_type = "pdf"
    elif ext in (".pptx", ".ppt"):
        slides = _parse_pptx(file_path)
        doc_type = "pptx"
    else:
        raise ValueError(f"Unsupported file type: {ext}. Supported: .pdf, .pptx")

    logger.info("Parsed %s (%s): %d slides/pages", filename, doc_type, len(slides))
    return {
        "filename": filename,
        "type": doc_type,
        "slides": slides,
    }
